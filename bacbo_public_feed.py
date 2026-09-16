"""Coleta Bac Bo do Casino.org e salva o historico em Excel.

Pipeline: API Casino.org -> requests -> pandas.DataFrame -> Excel.
"""

from __future__ import annotations

import argparse
import os
import sys
import time
from contextlib import contextmanager
from pathlib import Path

import pandas as pd
import requests
from openpyxl.styles import Font, PatternFill


API_URL = "https://api-cs.casino.org/svc-evolution-game-events/api/bacbo"
TABLE_ID = "BacBo00000000001"
BRASILIA_TZ = "America/Sao_Paulo"
COLUMNS = [
    "round_id", "table_id", "settled_at_brasilia", "result", "raw_result",
    "player_dice", "player_score", "banker_dice", "banker_score", "multiplier",
    "collected_at_brasilia", "source",
]


@contextmanager
def single_collector_lock(excel_path: Path):
    """Evita duas instancias gravando a mesma planilha."""
    lock_path = excel_path.with_suffix(excel_path.suffix + ".lock")
    lock_path.parent.mkdir(parents=True, exist_ok=True)
    handle = lock_path.open("a+b")
    handle.seek(0)
    handle.write(b"0")
    handle.flush()
    try:
        if os.name == "nt":
            import msvcrt

            handle.seek(0)
            msvcrt.locking(handle.fileno(), msvcrt.LK_NBLCK, 1)
        else:
            import fcntl

            fcntl.flock(handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
    except OSError as error:
        handle.close()
        raise RuntimeError(f"Outro coletor ja esta gravando {excel_path}.") from error
    try:
        yield
    finally:
        try:
            if os.name == "nt":
                handle.seek(0)
                msvcrt.locking(handle.fileno(), msvcrt.LK_UNLCK, 1)
            else:
                fcntl.flock(handle.fileno(), fcntl.LOCK_UN)
        finally:
            handle.close()


def result_name(outcome: object) -> tuple[str, str]:
    mapping = {"BankerWon": ("B", "banker"), "PlayerWon": ("P", "player"), "Tie": ("E", "tie")}
    try:
        return mapping[str(outcome)]
    except KeyError as error:
        raise ValueError(f"Resultado desconhecido: {outcome!r}") from error


def dice_values(dice: object) -> tuple[str, int | None]:
    if not isinstance(dice, dict):
        return "", None
    values = [dice.get("first"), dice.get("second")]
    return " + ".join(str(value) for value in values if value is not None), dice.get("score")


def fetch_dataframe(session: requests.Session, timeout: float, limit: int) -> pd.DataFrame:
    response = session.get(
        API_URL,
        params={"page": 0, "size": limit, "sort": "data.settledAt,desc", "duration": 1},
        timeout=timeout,
    )
    response.raise_for_status()
    payload = response.json()
    if not isinstance(payload, list):
        raise ValueError("A API nao retornou uma lista de rodadas.")

    rows: list[dict[str, object]] = []
    collected_at = pd.Timestamp.now(tz="UTC").tz_convert(BRASILIA_TZ).tz_localize(None)
    for item in payload:
        data = item.get("data", {}) if isinstance(item, dict) else {}
        result = data.get("result", {})
        table = data.get("table", {})
        if table.get("id") != TABLE_ID or not result or not data.get("settledAt"):
            continue
        short_result, raw_result = result_name(result.get("outcome"))
        player_dice, player_score = dice_values(result.get("playerDice"))
        banker_dice, banker_score = dice_values(result.get("bankerDice"))
        settled_at = pd.to_datetime(data["settledAt"], utc=True).tz_convert(BRASILIA_TZ).tz_localize(None)
        rows.append({
            "round_id": str(item.get("id") or data.get("id")),
            "table_id": TABLE_ID,
            "settled_at_brasilia": settled_at,
            "result": short_result,
            "raw_result": raw_result,
            "player_dice": player_dice,
            "player_score": player_score,
            "banker_dice": banker_dice,
            "banker_score": banker_score,
            "multiplier": result.get("multiplier"),
            "collected_at_brasilia": collected_at,
            "source": "casino.org",
        })
    return pd.DataFrame(rows, columns=COLUMNS)


def localize_datetime_column(series: pd.Series) -> pd.Series:
    """Aceita tanto o CSV legado quanto datas ja gravadas no Excel."""
    text = series.astype(str)
    legacy_mask = text.str.match(r"^\d{2}/\d{2}/\d{4} \d{2}:\d{2}:\d{2}")
    parsed = pd.Series(pd.NaT, index=series.index, dtype="datetime64[ns]")
    parsed.loc[legacy_mask] = pd.to_datetime(
        text.loc[legacy_mask].str.slice(0, 19),
        format="%d/%m/%Y %H:%M:%S",
        errors="coerce",
    )
    parsed.loc[~legacy_mask] = pd.to_datetime(text.loc[~legacy_mask], errors="coerce", format="mixed")
    return parsed


def normalize_dataframe(frame: pd.DataFrame) -> pd.DataFrame:
    normalized = frame.reindex(columns=COLUMNS).copy()
    if normalized.empty:
        return pd.DataFrame(columns=COLUMNS)
    for column in ("settled_at_brasilia", "collected_at_brasilia"):
        normalized[column] = localize_datetime_column(normalized[column])
    return normalized.dropna(subset=["round_id"]).drop_duplicates(subset=["round_id"], keep="last")


def load_history(excel_path: Path, legacy_csv: Path) -> pd.DataFrame:
    if excel_path.exists():
        return normalize_dataframe(pd.read_excel(excel_path, sheet_name="Resultados"))
    if legacy_csv.exists():
        return normalize_dataframe(pd.read_csv(legacy_csv))
    return pd.DataFrame(columns=COLUMNS)


def save_excel(frame: pd.DataFrame, excel_path: Path) -> None:
    excel_path.parent.mkdir(parents=True, exist_ok=True)
    ordered = normalize_dataframe(frame).sort_values("settled_at_brasilia", kind="stable")
    with pd.ExcelWriter(excel_path, engine="openpyxl", datetime_format="DD/MM/YYYY HH:MM:SS") as writer:
        ordered.to_excel(writer, sheet_name="Resultados", index=False)
        worksheet = writer.sheets["Resultados"]
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        worksheet.sheet_view.showGridLines = False
        header_fill = PatternFill("solid", fgColor="17365D")
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
        widths = {"A": 28, "B": 22, "C": 24, "D": 10, "E": 12, "F": 14, "G": 12, "H": 14, "I": 12, "J": 12, "K": 24, "L": 14}
        for column, width in widths.items():
            worksheet.column_dimensions[column].width = width
        for row in range(2, worksheet.max_row + 1):
            worksheet.cell(row, 3).number_format = 'DD/MM/YYYY HH:MM:SS "BRT"'
            worksheet.cell(row, 11).number_format = 'DD/MM/YYYY HH:MM:SS "BRT"'


def main() -> int:
    parser = argparse.ArgumentParser(description="Casino.org -> requests -> pandas -> Excel")
    parser.add_argument("--excel", type=Path, default=Path("bacbo_casinoscores_resultados.xlsx"))
    parser.add_argument("--legacy-csv", type=Path, default=Path("bacbo_casinoscores_results.csv"))
    parser.add_argument("--interval", type=float, default=35, help="segundos entre consultas")
    parser.add_argument("--timeout", type=float, default=15)
    parser.add_argument("--limit", type=int, default=10)
    parser.add_argument("--once", action="store_true")
    parser.add_argument("--rebuild-from-csv", action="store_true", help="recria o Excel a partir do CSV legado")
    args = parser.parse_args()
    if min(args.interval, args.timeout, args.limit) <= 0:
        parser.error("--interval, --timeout e --limit devem ser maiores que zero")

    session = requests.Session()
    session.trust_env = False
    session.headers.update({"User-Agent": "bacbo-casino-org-excel/1.0", "Accept": "application/json"})
    try:
        with single_collector_lock(args.excel):
            history = (
                normalize_dataframe(pd.read_csv(args.legacy_csv))
                if args.rebuild_from_csv and args.legacy_csv.exists()
                else load_history(args.excel, args.legacy_csv)
            )
            if args.rebuild_from_csv:
                save_excel(history, args.excel)
            while True:
                try:
                    latest = fetch_dataframe(session, args.timeout, args.limit)
                    updated = normalize_dataframe(pd.concat([history, latest], ignore_index=True))
                    if args.rebuild_from_csv or len(updated) != len(history) or not args.excel.exists():
                        save_excel(updated, args.excel)
                    history = updated
                except (requests.RequestException, ValueError, OSError) as error:
                    print(f"Falha na coleta: {error}", file=sys.stderr)
                if args.once:
                    return 0
                time.sleep(args.interval)
    except RuntimeError as error:
        print(error, file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
